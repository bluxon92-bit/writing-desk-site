# -*- coding: utf-8 -*-
import openpyxl, yaml, re, datetime as dt
SRC='recalc/writing-desk-hub-master-sheet_filled.xlsx'
OUT='/home/claude/hub-build/_data/hub/'
wb=openpyxl.load_workbook(SRC,data_only=True)
A=wb['Apps']
hdr={A.cell(4,c).value:c for c in range(1,A.max_column+1)}
def col(name): return hdr[name]
CHK=re.compile(r"\s*\[CHECK:[^\]]*\]")
todo=[]
def clean(s,where=''):
    if s is None: return None
    s=str(s).strip()
    if CHK.search(s): todo.append((where,CHK.findall(s)))
    s=CHK.sub('',s).strip()
    return s or None
SYM={'USD':'$','GBP':'£','EUR':'€'}
def money(v,cur):
    if v is None: return None
    v=float(v); n=('%d'%v) if v==int(v) else ('%.2f'%v)
    return f"{SYM.get(cur,'')}{n}"
UNIT={'/mo':'/mo','/yr':'/yr','one-time':' one-time'}
def amount(v,u,cur):
    m=money(v,cur)
    if m is None: return None
    return m+UNIT.get(u,'')
def price_line(r):
    cur=r['Currency']; ent=r['Entry price']; eu=r['Entry unit']; top=r['Top price']; tu=r['Top unit']
    if eu is None and ent is None: return None
    code=f" ({cur})" if cur in ('USD','EUR') else ''
    if eu=='free':
        s='Free'
        if top: s+=f"; paid plans up to {amount(top,tu,cur)}{code}"
        return s
    s=f"From {amount(ent,eu,cur)}{code}"
    if top and (top!=ent or tu!=eu) and not (eu=='one-time' and tu=='one-time'):
        s+=(f" or {amount(top,tu,cur)}" if tu!=eu else f"; up to {amount(top,tu,cur)}")
    return s
FEATS=[('Plotting / outlining','Plotting and outlining'),('Worldbuilding / story bible','Worldbuilding'),('Markdown','Markdown'),
 ('Export (EPUB / PDF / DOCX)','Export to EPUB, PDF or DOCX'),('Citations / references','Citations'),('Grammar / style checking','Grammar and style checks'),
 ('Collaboration','Collaboration'),('Offline / local-first','Works offline')]
PLAT=[('Mac','Mac'),('Browser','Browser'),('Windows','Windows'),('Linux','Linux'),('iOS','iOS'),('Android','Android')]
# find platform headers
ph={}
for c in range(1,A.max_column+1):
    v=A.cell(4,c).value
    if v in ('Mac','Browser','Windows','Linux','iOS','Android'): ph[v]=c
apps={}
order=[]
for rr in range(5,A.max_row+1):
    name=A.cell(rr,2).value
    if not name: continue
    r={h:A.cell(rr,c).value for h,c in hdr.items()}
    if r['Include on hub?']!='Y': continue
    slug=r['Slug']
    plats=[]
    for k,_ in PLAT:
        v=A.cell(rr,ph[k]).value
        if v=='Y': plats.append({'name':k})
        elif v=='Partial': plats.append({'name':k,'note':'limited'})
    feats=[]
    for k,label in FEATS:
        v=r[k]
        if v=='Y': feats.append(label)
        elif v=='Partial': feats.append(label+' (partial)')
    stance=r['AI stance']
    chip=None
    if stance=='None': chip='No AI'
    elif stance in ('Assistive','Analysis only','Generative','General assistant'): chip='AI: '+stance.lower() if stance!='Analysis only' else 'AI: analysis only'
    rating=None
    if r['Rating source']=='Capterra' and r['Rating (out of 5)']:
        d=r['Rating checked']
        rating={'source':'Capterra','value':r['Rating (out of 5)'],'checked':d.date().isoformat() if d else None}
    cmp=None
    if r['Full comparison page']=='Live' and r['Comparison page URL']:
        cmp=re.sub(r'^https?://[^/]+','',r['Comparison page URL'])
    chk=r['Price last checked']
    d=dict(name=name,slug=slug,status=r['Status'],category=r['Primary category'],
       tags=[t.strip() for t in (r['Tags'] or '').split(';') if t.strip()],
       overview=clean(r['Overview (what it is)'],name),best_for=clean(r['Best for (short)'],name),
       choose=clean(r['Choose it if...'],name),skip=clean(r['Skip it if...'],name),
       platforms=plats,mac_first=(r['Mac-first']=='Y'),
       pricing_model=r['Pricing model'],price_line=price_line(r),
       price_checked=chk.date().isoformat() if chk else None,
       free_tier=r['Free tier / trial'],features=feats,ai_chip=chip,rating=rating,
       comparison=cmp)
    if r['Logo (Cloudinary URL)']: d['logo']=r['Logo (Cloudinary URL)']
    apps[slug]=d; order.append(slug)
# Writing Desk: two currencies, stated on the card
apps['writing-desk']['price_line']="Free plan; Premium from £8/mo billed annually (£10 monthly) or $10/mo USD billed annually ($12 monthly)"
apps['scrivener']['price_line']='£55 one-time per desktop platform; £88 multi-device bundle; iOS $23.99 (USD)'
apps['writing-desk']['platforms']=[{'name':'Mac'}]
apps['writing-desk']['free_tier']="Free plan, plus a 30-day Premium trial"
# categories
CT=wb['Categories']; ch={CT.cell(4,c).value:c for c in range(1,CT.max_column+1)}
CL=wb['Category Lists']
lists={}
for rr in range(5,CL.max_row+1):
    cat,app,fl=[CL.cell(rr,k).value for k in (1,2,3)]
    if not app: continue
    slug=next((s for s,a in apps.items() if a['name']==app),None)
    if not slug: continue
    lists.setdefault(cat,{'F':[],'L':[]})[fl].append(slug)
cats={}
for rr in range(5,12):
    page=CT.cell(rr,1).value
    g=lambda k: clean(CT.cell(rr,ch[k]).value,page)
    url=g('URL'); key={'Hub index':'index','Mac list':'mac'}.get(page) or url.strip('/').split('/')[-1]
    c=dict(page=page,url=url,h1=g('H1 (draft)'),title=g('Title tag (draft)'),meta=g('Meta description (draft)'),
       for_whom=g('Who it is for (one line)'),pick=g('My pick'),why=g('Why (first person, 150-300 words)'),
       method=g('How I tested'),runner_ups=[x.strip() for x in (g('Runner-ups') or '').split(';') if x.strip()],
       primary_keyword=g('Primary keyword'))
    if page in lists:
        c['featured']=lists[page]['F']; c['listed']=lists[page]['L']
    cats[key]=c
# mac list
ML=wb['Mac List']
mac_rows=[]
for rr in range(5,ML.max_row+1):
    n=ML.cell(rr,1).value
    if not n or ML.cell(rr,3).value!='Y': continue
    slug=next((s for s,a in apps.items() if a['name']==n),None)
    if not slug: continue
    note=clean(ML.cell(rr,5).value,'Mac '+n)
    if note: note=re.sub(r'\s*Built with Tauri\.?','',note).strip()
    apps[slug]['mac_note']=note
    mac_rows.append(slug)
wd=['writing-desk']
mf=sorted([s for s in mac_rows if s!='writing-desk' and apps[s]['mac_first']],key=lambda s:apps[s]['name'].lower())
ot=sorted([s for s in mac_rows if s!='writing-desk' and not apps[s]['mac_first']],key=lambda s:apps[s]['name'].lower())
# featured = Writing Desk, mac-first apps, and the apps with a live comparison page
feat=[s for s in wd+mf+ot if s=='writing-desk' or apps[s]['mac_first'] or apps[s]['comparison']]
cats['mac']['featured']=feat; cats['mac']['listed']=[s for s in wd+mf+ot if s not in feat]
# copy blocks
PC=wb['Page Copy']
blocks={PC.cell(r,1).value:PC.cell(r,3).value for r in range(5,PC.max_row+1) if PC.cell(r,1).value}
settings=dict(verified_only=True,documents_page_live=False,last_updated='2026-10-11',
  disclosure_hub=blocks['Disclosure (hub)'],disclosure_category=blocks['Disclosure (category pages)'],
  how_we_chose=blocks['How we chose'],ratings_note=blocks['Ratings note'],price_fallback=blocks['Price line fallback'],
  mac_first_tooltip=blocks['Mac-first badge tooltip'])
def dump(o,f):
    with open(OUT+f,'w',encoding='utf-8') as fh:
        yaml.safe_dump(o,fh,allow_unicode=True,sort_keys=False,width=10000,default_flow_style=False)
dump(apps,'apps.yml'); dump(cats,'categories.yml'); dump(settings,'settings.yml')
print(len(apps),'apps;',sum(1 for a in apps.values() if a['status']=='Verified'),'verified')
for k,c in cats.items(): print(k,len(c.get('featured',[])),len(c.get('listed',[])))
print('TODO stripped:',todo)
import collections
print(collections.Counter(a['status'] for a in apps.values()))
print('no price:',[a['name'] for a in apps.values() if not a['price_line']])
print('no choose/skip/overview:',[a['name'] for a in apps.values() if not (a['choose'] and a['skip'] and a['overview'])])
